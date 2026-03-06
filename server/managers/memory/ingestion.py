import csv
import io
import os
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Any, List, Optional

from managers.memory.retrieval import RetrievalService


class FileStore:
    def save_bytes(self, filename: str, content: bytes) -> str:
        raise NotImplementedError


class LocalFileStore(FileStore):
    def __init__(self, base_dir: str = "data/uploads"):
        self.base_dir = Path(base_dir)
        self.base_dir.mkdir(parents=True, exist_ok=True)

    def save_bytes(self, filename: str, content: bytes) -> str:
        safe_name = Path(filename).name
        out = self.base_dir / safe_name
        out.write_bytes(content)
        return str(out)


class S3FileStore(FileStore):
    def save_bytes(self, filename: str, content: bytes) -> str:
        raise NotImplementedError("S3 backend placeholder. Plug boto3 client here.")


class MinioFileStore(FileStore):
    def save_bytes(self, filename: str, content: bytes) -> str:
        raise NotImplementedError("MinIO backend placeholder. Plug minio client here.")


@dataclass
class IngestionResult:
    doc_id: str
    source_name: str
    source_type: str
    chunks: int


class IngestionService:
    def __init__(self, retriever: RetrievalService, store: Optional[FileStore] = None):
        self.retriever = retriever
        self.store = store or LocalFileStore()

    def _extract_text(self, path: str) -> str:
        ext = Path(path).suffix.lower()
        if ext in {".md", ".txt", ".py", ".js", ".ts", ".json", ".yaml", ".yml", ".html", ".htm", ".css", ".java", ".go", ".rs", ".sql", ".sh"}:
            return Path(path).read_text(encoding="utf-8", errors="ignore")

        if ext == ".csv":
            rows = []
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    rows.append(", ".join(row))
            return "\n".join(rows)

        if ext == ".pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(path)
                return "\n".join((p.extract_text() or "") for p in reader.pages)
            except Exception:
                return ""

        if ext == ".docx":
            try:
                from docx import Document

                doc = Document(path)
                return "\n".join(p.text for p in doc.paragraphs)
            except Exception:
                return ""

        if ext in {".xlsx", ".xls"}:
            try:
                import openpyxl

                wb = openpyxl.load_workbook(path, read_only=True)
                texts = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        texts.append(", ".join("" if v is None else str(v) for v in row))
                return "\n".join(texts)
            except Exception:
                return ""

        if ext == ".zip":
            texts = []
            with zipfile.ZipFile(path) as zf:
                for name in zf.namelist():
                    lower = name.lower()
                    if lower.endswith((".py", ".js", ".ts", ".md", ".txt", ".java", ".go", ".rs", ".json", ".yaml", ".yml", ".html", ".css", ".sql")):
                        try:
                            texts.append(zf.read(name).decode("utf-8", errors="ignore"))
                        except Exception:
                            pass
            return "\n\n".join(texts)

        return Path(path).read_text(encoding="utf-8", errors="ignore")

    def _chunk_text(self, text: str, chunk_size: int = 900, overlap: int = 180) -> List[str]:
        text = (text or "").strip()
        if not text:
            return []

        chunks: List[str] = []
        start = 0
        n = len(text)
        while start < n:
            end = min(n, start + chunk_size)
            part = text[start:end]
            if part.strip():
                chunks.append(part.strip())
            if end == n:
                break
            start = max(start + 1, end - overlap)
        return chunks

    def ingest_bytes(
        self,
        filename: str,
        content: bytes,
        source_type: str = "upload",
        metadata: Optional[Dict[str, Any]] = None,
    ) -> IngestionResult:
        saved_path = self.store.save_bytes(filename, content)
        text = self._extract_text(saved_path)
        chunks = self._chunk_text(text)
        doc_id = self.retriever.index_document(
            source_name=Path(filename).name,
            source_type=source_type,
            chunks=chunks,
            metadata={**(metadata or {}), "stored_path": saved_path},
        )
        return IngestionResult(
            doc_id=doc_id,
            source_name=Path(filename).name,
            source_type=source_type,
            chunks=len(chunks),
        )
